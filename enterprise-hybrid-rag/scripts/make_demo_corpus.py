"""Generate the synthetic demo corpus used by the README, tests and evaluation.

The corpus describes a fictional company ("Northwind Analytics"). It is written
here as structured data and rendered to three real file formats so that the
loaders are exercised against genuine PDF/DOCX/Markdown bytes rather than
fixtures that happen to parse.

Run:  python scripts/make_demo_corpus.py
"""

from __future__ import annotations

import sys
import textwrap
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import docx
import fitz
from docx.shared import Pt

from app.config.settings import get_settings

Section = tuple[str, list[str]]

HANDBOOK: list[Section] = [
    ("Introduction", [
        "This handbook describes the employment policies of Northwind Analytics for all "
        "permanent employees. It applies to every office and to employees working remotely. "
        "Where local law conflicts with this handbook, local law prevails.",
        "The handbook is maintained by the People Operations team and is reviewed every "
        "twelve months. The current revision is HB-7.2, effective 1 January 2025.",
    ]),
    ("Working Hours", [
        "The standard working week is 40 hours, normally worked between Monday and Friday. "
        "Core collaboration hours are 10:00 to 16:00 in the employee's local time zone, "
        "during which employees are expected to be reachable.",
        "Employees may shift their start time between 07:00 and 10:00 with the agreement of "
        "their line manager. Overtime is not expected; where it occurs, time off in lieu is "
        "granted at a one to one ratio and must be taken within 60 days.",
    ]),
    ("Annual Leave", [
        "Employees are entitled to 22 days of paid annual leave per calendar year, in "
        "addition to public holidays observed in their country of employment.",
        "Annual leave accrues monthly at a rate of 1.83 days per completed month of service. "
        "Employees joining mid-year receive a pro-rated entitlement.",
        "A maximum of 5 unused annual leave days may be carried over into the following "
        "calendar year. Carried-over days expire on 31 March and are not paid out.",
        "Requests for more than three consecutive days of leave must be submitted at least "
        "10 working days in advance. Line managers must approve or decline a leave request "
        "within 5 working days of submission.",
    ]),
    ("Sick Leave", [
        "Employees receive 10 paid sick days per calendar year. Unused sick days do not "
        "carry over and are not paid out on termination.",
        "Employees must notify their line manager before 09:30 on the first day of absence. "
        "A medical certificate is required from the fourth consecutive day of absence.",
        "Absence beyond 10 days in a calendar year is handled under the long-term absence "
        "procedure, which is managed by People Operations together with the employee.",
    ]),
    ("Parental Leave", [
        "The primary caregiver is entitled to 16 weeks of parental leave at full pay. The "
        "secondary caregiver is entitled to 4 weeks at full pay.",
        "Parental leave must start within 12 months of the birth or adoption of the child "
        "and may be taken in up to three separate blocks with manager agreement.",
        "Employees returning from parental leave may request a reduced schedule of 80 "
        "percent for up to six months at 80 percent of salary.",
    ]),
    ("Remote Work", [
        "Employees may work remotely up to 3 days per week by default. Fully remote "
        "arrangements require written approval from a Vice President and are reviewed "
        "annually.",
        "A home office stipend of 400 US dollars is available every two years and covers "
        "desk, chair, monitor and lighting equipment. The stipend does not cover internet "
        "subscriptions or coworking memberships.",
        "Employees working from a country other than their country of employment for more "
        "than 30 days in a rolling twelve month period must obtain approval from People "
        "Operations because of tax and immigration exposure.",
    ]),
    ("Expenses and Travel", [
        "Expense claims must be submitted within 30 days of the date the expense was "
        "incurred. Claims submitted after 60 days will not be reimbursed.",
        "The meal per diem is 65 US dollars for domestic travel and 95 US dollars for "
        "international travel. Receipts are not required for per diem claims.",
        "Flights under 6 hours must be booked in economy class. Business class may be "
        "booked for flights longer than 8 hours with prior manager approval.",
        "Hotel spending is capped at 220 US dollars per night for domestic travel and 300 "
        "US dollars per night for international travel, excluding local taxes.",
    ]),
    ("Performance and Promotion", [
        "Performance reviews take place twice per year, in April and in October. Reviews "
        "use a five point rating scale where 3 means fully meets expectations.",
        "Promotion decisions are made only in the October cycle. A promotion case requires "
        "a written submission from the line manager and two peer references.",
        "Employees rated 1 are placed on a structured 60 day performance improvement plan "
        "with fortnightly checkpoints.",
    ]),
    ("Referrals", [
        "Employees receive a referral bonus of 1500 US dollars when a referred candidate is "
        "hired and completes 6 months of continuous service. There is no limit on the "
        "number of referrals an employee may make.",
    ]),
    ("Probation and Notice", [
        "New employees serve a probation period of 3 months, which may be extended once by "
        "up to 3 additional months.",
        "The notice period is 4 weeks during the first year of service, 6 weeks from the "
        "second year, and 8 weeks after two completed years of service.",
        "On termination, accrued and unused annual leave from the current calendar year is "
        "paid out. Carried-over days from a previous year are not paid out.",
    ]),
]

SECURITY: list[Section] = [
    ("Scope and Ownership", [
        "This Information Security Policy, reference ISP-2024-03, applies to all employees, "
        "contractors and third parties who access Northwind Analytics systems or data.",
        "The policy is owned by the Chief Information Security Officer and is formally "
        "reviewed annually. Exceptions require a documented risk acceptance signed by the "
        "CISO and are valid for a maximum of 90 days.",
    ]),
    ("Access Control", [
        "Access is granted on the principle of least privilege. Standard access requests are "
        "approved by the system owner and fulfilled within 2 business days.",
        "User access reviews are performed quarterly. Privileged and administrator access is "
        "reviewed monthly.",
        "Just-in-time elevation to production is limited to a maximum of 8 hours per session "
        "and requires a linked incident or change ticket.",
        "Accounts of departing employees are disabled within 4 hours of their final working "
        "hour and deleted after 30 days.",
    ]),
    ("Passwords and Multi-Factor Authentication", [
        "Passwords must be at least 14 characters long. Scheduled password rotation is not "
        "required; passwords are rotated only on suspected compromise, in line with current "
        "NIST guidance.",
        "Multi-factor authentication is mandatory for all production systems, for the VPN, "
        "and for the single sign-on provider. Hardware security keys are required for "
        "administrators.",
        "Shared accounts are prohibited. Where a shared service account is unavoidable, "
        "credentials must be stored in the approved secrets manager and rotated every 90 "
        "days.",
    ]),
    ("Data Classification", [
        "Data is classified into four levels: Public, Internal, Confidential and Restricted.",
        "Restricted data includes customer personal data, payment data and authentication "
        "secrets. Restricted data must be encrypted at rest using AES-256 and in transit "
        "using TLS 1.3.",
        "Confidential data may not be copied to personal devices or personal cloud storage "
        "accounts under any circumstances.",
        "All company laptops must use full disk encryption and an automatic screen lock "
        "after 10 minutes of inactivity.",
    ]),
    ("Incident Response", [
        "Security incidents are classified from SEV-1, the most severe, to SEV-4. A SEV-1 "
        "incident must be acknowledged by the on-call security engineer within 15 minutes.",
        "Suspected security incidents must be reported to the security team within 1 hour of "
        "discovery using the security incident channel.",
        "Where an incident involves personal data, affected customers and the relevant "
        "supervisory authority are notified within 72 hours of confirmation.",
        "A written post-incident review is completed within 10 business days for every SEV-1 "
        "and SEV-2 incident.",
    ]),
    ("Security Training", [
        "All employees complete mandatory security awareness training annually. The training "
        "takes approximately 45 minutes.",
        "Phishing simulations are run quarterly. Employees who fail two consecutive "
        "simulations are enrolled in additional targeted training.",
    ]),
    ("Vendor and Third-Party Risk", [
        "Vendors that process Restricted data must provide a current SOC 2 Type II report or "
        "an equivalent ISO 27001 certification before contract signature.",
        "Vendor security assessments are repeated every 12 months, or immediately following "
        "a publicly disclosed breach affecting that vendor.",
    ]),
    ("Logging and Audit", [
        "Security relevant logs are retained for 400 days. Audit logs are written to "
        "append-only storage and cannot be modified by application teams.",
        "Production access is logged and reviewed weekly by the security operations team.",
    ]),
]

ONBOARDING = """# Engineering Onboarding Guide

Welcome to the Northwind Analytics engineering organisation. This guide covers your first
two weeks. It is maintained by the Developer Experience team.

## Week One Checklist

On day one you receive your laptop, single sign-on credentials and a hardware security key.
Your buddy is assigned automatically and will meet you daily for the first week.

By the end of week one you should have shipped one small pull request to production. This is
deliberate: the onboarding target is a first merged pull request within 5 working days.

## Development Environment

All services target Python 3.11 and Node 20. Docker Desktop or an equivalent container
runtime is required.

Run `make bootstrap` in the monorepo root. It installs dependencies, seeds a local Postgres
database and starts the supporting containers. A full bootstrap takes about 12 minutes on a
current laptop.

Local secrets are read from `.env.local`, which is never committed. Production secrets live
in the secrets manager and are never available on developer machines.

## Git Workflow

We use trunk-based development. The `main` branch is protected and always deployable.

Branch names follow the pattern `feat/ENG-1234-short-description`, using `feat/`, `fix/` or
`chore/` followed by the ticket identifier. Merges into `main` are squash merges.

## Code Review

Pull requests should change no more than 400 lines. Larger changes must be split unless the
change is a mechanical refactor.

Service code requires 2 approvals before merge. Documentation-only changes require 1
approval. The author must not approve their own pull request.

Reviewers are expected to give a first response within 1 business day.

## Continuous Integration and Deployment

The pipeline runs five stages in order: lint, unit tests, integration tests, build, and
deploy to staging. Deployment to production requires a manual approval from a service owner.

New code must maintain at least 80 percent line coverage. The pipeline fails below that
threshold.

Staging deployments happen automatically on every merge to `main`. Production deployments
are batched and released twice per week, on Tuesday and Thursday.

## On-Call

Each service team runs a weekly on-call rotation managed in PagerDuty. The handover happens
every Wednesday at 14:00 UTC.

The acknowledgement target for a paging alert is 5 minutes. Every alert must link to a
runbook stored under the `/runbooks` directory of the service repository.

On-call engineers receive a compensation allowance of 250 US dollars per completed week of
rotation.

## Platform Architecture

The platform consists of four primary services: `ingestion-api`, `feature-store`,
`model-gateway` and `dashboard-web`.

Persistent state lives in Postgres 15. Redis is used for caching and rate limiting. Kafka
carries event streams between the ingestion API and the feature store.

The model gateway is the only component permitted to call external model providers. Direct
provider calls from other services are blocked at the network layer.

## Glossary

Feature drift refers to a change in the statistical distribution of a model input relative
to its training distribution. Drift is monitored daily and alerts at a population stability
index above 0.2.
"""

SLA = """# Customer Support Service Level Agreement

This document defines the support commitments of Northwind Analytics for customers on the
Standard and Enterprise plans. It is contractually binding and is reviewed each quarter.

## Coverage Hours

Standard plan support operates 08:00 to 20:00 Central European Time, Monday to Friday,
excluding public holidays in Germany.

Enterprise plan support operates 24 hours a day, 7 days a week, including public holidays.

## Priority Definitions

A P1 issue is a complete production outage or a confirmed data loss event affecting multiple
users.

A P2 issue is a severe degradation where a major feature is unusable but a workaround exists.

A P3 issue is a limited problem affecting a single user or a non-critical feature.

A P4 issue is a question, a documentation problem or a cosmetic defect.

## Response and Resolution Targets

The first response target is 30 minutes for P1, 4 business hours for P2, 1 business day for
P3 and 3 business days for P4.

The resolution target for a P1 issue is 8 hours. During an active P1, the support team
provides a written status update every 60 minutes until the issue is resolved.

Response targets are measured from the time a ticket is created in the support portal, not
from the time an email is sent.

## Escalation

If two consecutive status updates are missed during a P1, the customer may escalate directly
to the duty manager.

Any P1 that remains unresolved after 4 hours is automatically escalated to the Vice
President of Customer Support.

## Availability Commitment and Service Credits

The monthly uptime commitment is 99.9 percent for the Standard plan and 99.95 percent for
the Enterprise plan. Uptime is measured per calendar month.

If monthly uptime falls below the commitment but remains at or above 99.0 percent, the
customer receives a service credit of 10 percent of the monthly fee.

If monthly uptime falls below 99.0 percent but remains at or above 95.0 percent, the service
credit is 25 percent of the monthly fee.

If monthly uptime falls below 95.0 percent, the service credit is 50 percent of the monthly
fee.

Service credits must be claimed in writing within 30 days of the end of the affected month.
Credits are applied to the next invoice and are not refundable in cash.

## Maintenance Windows

Planned maintenance takes place on Sundays between 02:00 and 06:00 Central European Time.
Customers receive at least 7 days notice of planned maintenance.

Emergency maintenance may be performed at any time with a minimum of 4 hours notice, except
where an immediate security fix is required.

Planned maintenance windows are excluded from the uptime calculation. Emergency maintenance
is included.
"""


def render_pdf(path: Path, title: str, sections: list[Section]) -> None:
    """Lay out the handbook as a real multi-page PDF with heading font sizes."""
    doc = fitz.open()
    page = doc.new_page()
    margin, width, bottom = 62, 472, 782
    # A page cursor advanced by fractional line heights, not an integer.
    y: float = margin

    def new_page() -> None:
        nonlocal page, y
        page = doc.new_page()
        y = margin

    def write(text: str, size: float, font: str, gap: float) -> None:
        nonlocal y
        wrap = int(width / (size * 0.50))
        for line in textwrap.wrap(text, width=wrap) or [""]:
            if y > bottom:
                new_page()
            page.insert_text((margin, y), line, fontsize=size, fontname=font)
            y += size * 1.38
        y += float(gap)

    write(title, 18, "hebo", 14)
    for heading, paragraphs in sections:
        if y > bottom - 80:
            new_page()
        write(heading, 14, "hebo", 6)
        for paragraph in paragraphs:
            write(paragraph, 10.5, "helv", 8)

    doc.set_metadata({"title": title, "author": "Northwind Analytics"})
    doc.save(str(path))
    doc.close()


def render_docx(path: Path, title: str, sections: list[Section]) -> None:
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    document.add_heading(title, level=0)
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    document.core_properties.title = title
    document.save(str(path))


def main() -> None:
    from demo_corpus_part2 import (
        CONTRACTOR_MD,
        EXTRA_HANDBOOK,
        EXTRA_SECURITY,
        PROCUREMENT,
        RELEASE_MD,
        RETENTION_MD,
    )
    from demo_corpus_part3 import (
        HR_RECRUITMENT_MD,
        IT_SERVICE_DESK_MD,
        MODEL_GOVERNANCE_MD,
        OFFICE_FACILITIES_MD,
    )

    settings = get_settings()
    raw = Path(settings.raw_dir)
    raw.mkdir(parents=True, exist_ok=True)

    render_pdf(raw / "employee_handbook.pdf", "Northwind Analytics Employee Handbook",
               HANDBOOK + EXTRA_HANDBOOK)
    render_pdf(raw / "finance_procurement_policy.pdf",
               "Northwind Analytics Finance and Procurement Policy", PROCUREMENT)
    render_docx(raw / "information_security_policy.docx",
                "Northwind Analytics Information Security Policy", SECURITY + EXTRA_SECURITY)
    (raw / "engineering_onboarding.md").write_text(ONBOARDING, encoding="utf-8")
    (raw / "customer_support_sla.md").write_text(SLA, encoding="utf-8")
    (raw / "contractor_engagement_policy.md").write_text(CONTRACTOR_MD, encoding="utf-8")
    (raw / "data_retention_schedule.md").write_text(RETENTION_MD, encoding="utf-8")
    (raw / "product_release_process.md").write_text(RELEASE_MD, encoding="utf-8")
    (raw / "it_service_desk_guide.md").write_text(IT_SERVICE_DESK_MD, encoding="utf-8")
    (raw / "hr_recruitment_process.md").write_text(HR_RECRUITMENT_MD, encoding="utf-8")
    (raw / "model_governance_policy.md").write_text(MODEL_GOVERNANCE_MD, encoding="utf-8")
    (raw / "office_facilities_guide.md").write_text(OFFICE_FACILITIES_MD, encoding="utf-8")

    from demo_corpus_part4 import BCP_MD, DPA_MD, INCIDENT_MD, REMOTE_MD, TRAVEL_MD, VENDOR_MD
    from demo_corpus_part5 import CODE_OF_CONDUCT, COMPENSATION, DATA_QUALITY_MD, ML_PLATFORM_MD

    (raw / "travel_expense_policy.md").write_text(TRAVEL_MD, encoding="utf-8")
    (raw / "remote_work_policy.md").write_text(REMOTE_MD, encoding="utf-8")
    (raw / "incident_response_runbook.md").write_text(INCIDENT_MD, encoding="utf-8")
    (raw / "business_continuity_plan.md").write_text(BCP_MD, encoding="utf-8")
    (raw / "data_processing_agreement.md").write_text(DPA_MD, encoding="utf-8")
    (raw / "vendor_risk_assessment.md").write_text(VENDOR_MD, encoding="utf-8")
    (raw / "ml_platform_standards.md").write_text(ML_PLATFORM_MD, encoding="utf-8")
    (raw / "data_engineering_standards.md").write_text(DATA_QUALITY_MD, encoding="utf-8")
    render_pdf(raw / "compensation_and_benefits.pdf",
               "Northwind Analytics Compensation and Benefits Policy", COMPENSATION)
    render_docx(raw / "code_of_conduct.docx",
                "Northwind Analytics Code of Conduct", CODE_OF_CONDUCT)

    for file in sorted(raw.iterdir()):
        print(f"{file.name:42} {file.stat().st_size / 1024:7.1f} KB")


if __name__ == "__main__":
    main()
